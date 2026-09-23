from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Polygon
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "src"))

from argus_morphing.airfoil_io import (  # noqa: E402
    cosine_grid,
    interpolate_surfaces,
    read_openvsp_airfoil_csv,
)
from argus_morphing.airfoil_morphing import morph_airfoil  # noqa: E402


BLUE = "#174A6B"
TEAL = "#237A72"
GREEN = "#4E8B57"
ORANGE = "#C56A32"
RED = "#B4473F"
PURPLE = "#785B9E"
LIGHT_BLUE = "#DCE8EF"
LIGHT_GREEN = "#DDEADF"
GRAY = "#5B6670"


def read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def number(row, name, default=0.0):
    try:
        return float(row.get(name, default))
    except (TypeError, ValueError):
        return default


def save_figure(fig, output: Path, name: str):
    for suffix in [".svg", ".pdf"]:
        fig.savefig(output / f"{name}{suffix}", bbox_inches="tight")
    fig.savefig(output / f"{name}.png", dpi=240, bbox_inches="tight")
    plt.close(fig)


def station_geometry(project: Path):
    rows = read_csv(project / "outputs" / "refined_baseline" / "refined_section_table.csv")
    y = [number(rows[0], "y_start")]
    chord = [number(rows[0], "chord")]
    xle = [0.0]
    twist = [number(rows[0], "twist_increment_deg")]
    for row in rows[1:]:
        dy = number(row, "span")
        y.append(number(row, "y_end"))
        xle.append(xle[-1] + np.tan(np.radians(number(row, "sweep_deg"))) * dy)
        chord.append(number(row, "chord"))
        twist.append(number(row, "twist_increment_deg"))
    return np.asarray(y), np.asarray(xle), np.asarray(chord), np.asarray(twist)


def overview_figure(project: Path, output: Path):
    y, xle, chord, twist = station_geometry(project)
    eta = y / y[-1]
    xte = xle + chord
    morph = (eta >= 0.60) & (eta <= 0.95)
    xhinge = xle + 0.62 * chord

    fig = plt.figure(figsize=(14, 7.5), constrained_layout=True)
    grid = fig.add_gridspec(2, 2, width_ratios=[1.65, 1.0], height_ratios=[1.0, 0.65])
    ax = fig.add_subplot(grid[:, 0])
    ax2 = fig.add_subplot(grid[0, 1])
    ax3 = fig.add_subplot(grid[1, 1])

    polygon = np.column_stack([
        np.r_[xle, xte[::-1]],
        np.r_[y, y[::-1]],
    ])
    ax.add_patch(Polygon(polygon, facecolor=LIGHT_BLUE, edgecolor=BLUE, lw=1.8))
    morph_polygon = np.column_stack([
        np.r_[xle[morph], xte[morph][::-1]],
        np.r_[y[morph], y[morph][::-1]],
    ])
    ax.add_patch(Polygon(morph_polygon, facecolor=LIGHT_GREEN, edgecolor=GREEN, lw=1.8))
    ax.plot(xhinge[morph], y[morph], color=ORANGE, lw=2.2, ls="--", label="Morphing start line, x_h/c=0.62")
    ax.plot(xle, y, color=BLUE, lw=1.4)
    ax.plot(xte, y, color=BLUE, lw=1.4)
    for index in range(len(y)):
        if morph[index]:
            ax.plot([xle[index], xte[index]], [y[index], y[index]], color=GREEN, lw=0.55, alpha=0.7)
    ax.axhline(0.60 * y[-1], color=GRAY, ls=":", lw=1.2)
    ax.axhline(0.95 * y[-1], color=GRAY, ls=":", lw=1.2)
    ax.text(
        3.22,
        0.60 * y[-1] + 0.08,
        "eta=0.60",
        va="center",
        fontsize=9,
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.8},
    )
    ax.text(
        3.48,
        0.95 * y[-1] + 0.08,
        "eta=0.95",
        va="center",
        fontsize=9,
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.8},
    )
    ax.text(0.15, 1.65, "Fixed wing region", color=BLUE, weight="bold", fontsize=12)
    ax.text(2.95, 4.55, "Morphing outer wing", color=GREEN, weight="bold", fontsize=12, rotation=28)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel("Wing-local x [m]")
    ax.set_ylabel("Semispan y [m]")
    ax.set_title("NASA EET AR=12 cruise wing and selected morphing region", fontsize=15)
    ax.grid(alpha=0.18)
    ax.legend(loc="lower right", fontsize=9)

    raw = project / "outputs" / "baseline" / "raw" / "baseline_airfoil_points.csv"
    upper, lower = read_openvsp_airfoil_csv(raw, 4)
    x = cosine_grid(201)
    upper_z, lower_z = interpolate_surfaces(upper, lower, x)
    morphed = morph_airfoil(x, upper_z, lower_z, 0.62, 0.02)
    ax2.fill_between(x, lower_z, upper_z, color=LIGHT_BLUE, alpha=0.8, label="Baseline section")
    ax2.plot(x, upper_z, color=BLUE, lw=1.5)
    ax2.plot(x, lower_z, color=BLUE, lw=1.5)
    ax2.plot(x, morphed.upper_z, color=GREEN, lw=2.0, label="Morphed, A/c=+0.02")
    ax2.plot(x, morphed.lower_z, color=GREEN, lw=2.0)
    ax2.axvline(0.62, color=ORANGE, ls="--", lw=1.7, label="x_h/c=0.62")
    ax2.annotate(
        "Trailing edge down",
        xy=(0.98, morphed.upper_z[-1]),
        xytext=(0.73, -0.088),
        arrowprops={"arrowstyle": "->", "color": RED},
        color=RED,
        fontsize=9,
    )
    ax2.set_aspect("equal", adjustable="box")
    ax2.set_ylim(-0.105, 0.075)
    ax2.set_xlabel("x/c")
    ax2.set_ylabel("z/c")
    ax2.set_title("Representative airfoil section, eta about 0.71")
    ax2.grid(alpha=0.22)
    ax2.legend(fontsize=7.5, loc="upper left", ncol=3)

    ax3.plot(eta, chord, color=BLUE, lw=2.2, marker="o", ms=3, label="Chord")
    ax3_t = ax3.twinx()
    ax3_t.plot(eta, twist, color=PURPLE, lw=2.0, marker="s", ms=3, label="Twist")
    ax3.axvspan(0.60, 0.95, color=LIGHT_GREEN, alpha=0.65)
    ax3.set_xlabel("eta = y/(b/2)")
    ax3.set_ylabel("Chord [m]", color=BLUE)
    ax3_t.set_ylabel("Twist [deg]", color=PURPLE)
    ax3.set_title("Baseline spanwise geometry")
    ax3.grid(alpha=0.22)

    fig.suptitle(
        "ARGUS baseline and morphing-wing definition",
        fontsize=18,
        weight="bold",
    )
    save_figure(fig, output, "00_baseline_morphing_overview")


def airfoil_figure(project: Path, output: Path):
    raw = project / "outputs" / "baseline" / "raw" / "baseline_airfoil_points.csv"
    upper, lower = read_openvsp_airfoil_csv(raw, 4)
    x = cosine_grid(201)
    upper_z, lower_z = interpolate_surfaces(upper, lower, x)
    cases = [(0.0, "Baseline", BLUE), (0.01, "A/c=+0.01", TEAL), (0.02, "A/c=+0.02", GREEN), (0.04, "A/c=+0.04", RED)]

    fig, axes = plt.subplots(2, 1, figsize=(11, 8), constrained_layout=True)
    for amplitude, label, color in cases:
        result = morph_airfoil(x, upper_z, lower_z, 0.62, amplitude)
        axes[0].plot(x, result.upper_z, color=color, lw=2, label=label)
        axes[0].plot(x, result.lower_z, color=color, lw=2)
        axes[1].plot(x, result.displacement, color=color, lw=2, label=label)
    axes[0].axvline(0.62, color=ORANGE, ls="--", lw=1.8)
    axes[0].text(0.63, 0.075, "Deformable region", color=ORANGE, fontsize=11)
    axes[0].set_aspect("equal", adjustable="box")
    axes[0].set_ylabel("z/c")
    axes[0].set_title("Thickness-preserving continuous camber morphing")
    axes[0].grid(alpha=0.22)
    axes[0].legend(ncol=4, fontsize=9)
    axes[1].axvline(0.62, color=ORANGE, ls="--", lw=1.8)
    axes[1].set_xlabel("x/c")
    axes[1].set_ylabel("Camber displacement, Delta z/c")
    axes[1].set_title("Smoothstep displacement: zero displacement and slope at x_h/c")
    axes[1].grid(alpha=0.22)
    save_figure(fig, output, "01_airfoil_morphing_definition")


def exact_load_figures(project: Path, output: Path):
    root = project / "outputs" / "exact_trim_loads"
    summary = read_csv(root / "exact_trim_load_summary.csv")
    selected_ids = ["refined_baseline", "doe_xh_0p62_A_p0p02_uniform"]
    labels = {"refined_baseline": "Refined baseline", selected_ids[1]: "Selected morphing candidate"}
    colors = {"refined_baseline": BLUE, selected_ids[1]: GREEN}
    styles = {"refined_baseline": "--", selected_ids[1]: "-"}
    loads = {
        case_id: read_csv(root / f"{case_id}_spanwise_load.csv")
        for case_id in selected_ids
    }

    fig, axes = plt.subplots(2, 1, figsize=(10.5, 8), sharex=True, constrained_layout=True)
    for case_id in selected_ids:
        rows = loads[case_id]
        eta = [number(row, "eta") for row in rows]
        axes[0].plot(eta, [number(row, "lift_per_span_N_per_m") for row in rows], color=colors[case_id], ls=styles[case_id], lw=2.4, label=labels[case_id])
        axes[1].plot(eta, [number(row, "sectional_cl") for row in rows], color=colors[case_id], ls=styles[case_id], lw=2.4, label=labels[case_id])
    for ax in axes:
        ax.axvspan(0.60, 0.95, color=LIGHT_GREEN, alpha=0.55)
        ax.grid(alpha=0.22)
        ax.legend(fontsize=9)
    axes[0].set_ylabel("Lift per span [N/m]")
    axes[0].set_title("Fixed-CL lift redistribution")
    axes[1].set_xlabel("eta = y/(b/2)")
    axes[1].set_ylabel("Sectional cl")
    axes[1].set_title("Sectional lift coefficient")
    save_figure(fig, output, "02_baseline_vs_selected_spanwise_loads")

    fig, axes = plt.subplots(3, 1, figsize=(10.5, 10), sharex=True, constrained_layout=True)
    for case_id in selected_ids:
        rows = loads[case_id]
        eta = np.asarray([number(row, "eta") for row in rows])
        dy = np.asarray([number(row, "dy_m") for row in rows])
        lift = np.asarray([number(row, "lift_per_span_N_per_m") for row in rows])
        y = np.asarray([number(row, "y_m") for row in rows])
        hinge = np.asarray([number(row, "hinge_moment_proxy_N") for row in rows])
        shear_outboard = np.cumsum((lift * dy)[::-1])[::-1]
        bend_outboard = np.asarray([
            np.sum(lift[index:] * dy[index:] * (y[index:] - y[index]))
            for index in range(len(y))
        ])
        axes[0].plot(eta, shear_outboard, color=colors[case_id], ls=styles[case_id], lw=2.3, label=labels[case_id])
        axes[1].plot(eta, bend_outboard, color=colors[case_id], ls=styles[case_id], lw=2.3)
        axes[2].plot(eta, hinge, color=colors[case_id], ls=styles[case_id], lw=2.3)
    for ax in axes:
        ax.axvspan(0.60, 0.95, color=LIGHT_GREEN, alpha=0.55)
        ax.grid(alpha=0.22)
    axes[0].set_ylabel("Outboard shear [N]")
    axes[0].set_title("Integrated aerodynamic loads toward the wing root")
    axes[0].legend(fontsize=9)
    axes[1].set_ylabel("Bending moment [N m]")
    axes[2].set_ylabel("Hinge-moment proxy [N]")
    axes[2].set_xlabel("eta = y/(b/2)")
    save_figure(fig, output, "03_selected_candidate_integrated_loads")

    return summary


def doe_figures(project: Path, output: Path, exact_summary: list[dict]):
    doe = read_csv(project / "outputs" / "doe" / "minimal_doe_ranked.csv")
    shapes = ["uniform", "bell", "tip_decreasing", "tip_increasing"]
    colors = dict(zip(shapes, [BLUE, PURPLE, TEAL, ORANGE]))
    fig, axes = plt.subplots(1, 3, figsize=(15, 4.8), sharey=True, constrained_layout=True)
    for ax, hinge in zip(axes, [0.35, 0.50, 0.62]):
        for shape in shapes:
            rows = [
                row for row in doe
                if abs(number(row, "x_h_over_c") - hinge) < 1e-8 and row["shape_type"] == shape
            ]
            rows.sort(key=lambda row: number(row, "A_max_over_c"))
            ax.plot(
                [number(row, "A_max_over_c") for row in rows],
                [-number(row, "CDi_change_percent") for row in rows],
                marker="o",
                lw=1.8,
                color=colors[shape],
                label=shape.replace("_", " "),
            )
        ax.axhline(0, color=GRAY, lw=1, ls=":")
        ax.set_title(f"x_h/c={hinge:.2f}")
        ax.set_xlabel("A_max/c")
        ax.grid(alpha=0.22)
    axes[0].set_ylabel("Induced-drag reduction [%]")
    axes[-1].legend(fontsize=8, loc="center left", bbox_to_anchor=(1.02, 0.5))
    save_figure(fig, output, "04_doe_induced_drag_response")

    candidates = [row for row in exact_summary if row["case_id"] != "refined_baseline"]
    fig, ax = plt.subplots(figsize=(8.8, 5.5))
    for row in candidates:
        x = number(row, "root_bending_change_percent_vs_baseline")
        y = -number(row, "CDi_change_percent_vs_baseline")
        rank = int(number(row, "rank"))
        color = GREEN if row["case_id"] == "doe_xh_0p62_A_p0p02_uniform" else BLUE
        ax.scatter(x, y, s=100, color=color, zorder=3)
        ax.annotate(f"#{rank}", (x, y), xytext=(7, 5), textcoords="offset points", fontsize=10)
    ax.set_xlabel("Half-wing root bending increase [%]")
    ax.set_ylabel("Induced-drag reduction [%]")
    ax.set_title("Aerodynamic benefit versus structural load penalty")
    ax.grid(alpha=0.22)
    ax.annotate(
        "Selected candidate\nx_h/c=0.62, A/c=+0.02",
        xy=(5.474, 4.334),
        xytext=(7.3, 4.32),
        arrowprops={"arrowstyle": "->", "color": GREEN},
        color=GREEN,
        weight="bold",
    )
    save_figure(fig, output, "05_aero_structural_tradeoff")

    baseline = next(row for row in exact_summary if row["case_id"] == "refined_baseline")
    chosen = next(row for row in exact_summary if row["case_id"] == "doe_xh_0p62_A_p0p02_uniform")
    fig, axes = plt.subplots(1, 3, figsize=(14.5, 5.2), constrained_layout=True)
    axes[0].bar(
        ["Baseline", "Selected"],
        [number(baseline, "CDi"), number(chosen, "CDi")],
        color=[BLUE, GREEN],
    )
    axes[0].set_ylabel("CDi")
    axes[0].set_title("Exact fixed-CL induced drag")
    axes[0].grid(axis="y", alpha=0.22)
    percent_names = ["CDi reduction", "Root bending increase"]
    percent_values = [
        -number(chosen, "CDi_change_percent_vs_baseline"),
        number(chosen, "root_bending_change_percent_vs_baseline"),
    ]
    bars = axes[1].bar(percent_names, percent_values, color=[GREEN, ORANGE])
    axes[1].bar_label(bars, labels=[f"{value:.2f}%" for value in percent_values], padding=4)
    axes[1].set_ylabel("Change versus baseline [%]")
    axes[1].set_title("Benefit and load penalty")
    axes[1].tick_params(axis="x", rotation=15)
    axes[1].grid(axis="y", alpha=0.22)
    load_names = ["Morph lift", "Bending", "Torque proxy"]
    load_values = [
        number(chosen, "morph_region_lift_N"),
        number(chosen, "morph_region_bending_about_eta_start_Nm"),
        abs(number(chosen, "morph_region_hinge_torque_proxy_Nm")),
    ]
    bars = axes[2].bar(load_names, load_values, color=[TEAL, PURPLE, RED])
    axes[2].bar_label(
        bars,
        labels=[f"{load_values[0]:.0f} N", f"{load_values[1]:.0f} N m", f"{load_values[2]:.0f} N m"],
        padding=4,
        fontsize=9,
    )
    axes[2].set_title("Morphing-region loads")
    axes[2].tick_params(axis="x", rotation=15)
    axes[2].grid(axis="y", alpha=0.22)
    save_figure(fig, output, "06_selected_candidate_summary")


def create_pptx(output: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slides = [
        ("Baseline and morphing-wing definition", "00_baseline_morphing_overview.png"),
        ("Continuous airfoil morphing definition", "01_airfoil_morphing_definition.png"),
        ("Fixed-lift spanwise aerodynamic loads", "02_baseline_vs_selected_spanwise_loads.png"),
        ("Integrated loads for actuator concept work", "03_selected_candidate_integrated_loads.png"),
        ("DOE induced-drag response", "04_doe_induced_drag_response.png"),
        ("Aerodynamic benefit and structural penalty", "05_aero_structural_tradeoff.png"),
        ("Selected candidate summary", "06_selected_candidate_summary.png"),
    ]
    for title, image in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.16), Inches(12.4), Inches(0.55))
        paragraph = title_box.text_frame.paragraphs[0]
        paragraph.text = title
        paragraph.font.name = "Aptos Display"
        paragraph.font.size = Pt(24)
        paragraph.font.bold = True
        paragraph.alignment = PP_ALIGN.LEFT
        slide.shapes.add_picture(str(output / image), Inches(0.45), Inches(0.82), width=Inches(12.4), height=Inches(6.3))
    prs.save(output / "ARGUS_editable_plot_deck.pptx")


def write_readme(output: Path):
    text = """# ARGUS editable plot package

All scientific plots are provided as:

- SVG: editable vector text, lines, colors, markers, and layout.
- PDF: vector format for publication and LaTeX workflows.
- PNG: convenient preview and direct PowerPoint use.
- PPTX: one slide per figure for rapid presentation assembly.

The SVG files can be edited in Inkscape, Adobe Illustrator, or PowerPoint
after inserting the SVG and converting it to shapes. The plotting source is
`scripts/13_create_editable_project_plots.py`.

Load convention:

- Wing-only VSPAERO model.
- Fixed target CL = 0.428277635.
- Dimensional loads use rho = 1.225 kg/m^3 and V = 45 m/s.
- Morphing region: eta = 0.60 to 0.95.
- Selected screening candidate: x_h/c = 0.62, A_max/c = +0.02, uniform spanwise amplitude.
- The hinge-moment quantity is a conceptual aerodynamic proxy, not a final
  actuator sizing torque.
"""
    (output / "README_editable_plots.md").write_text(text, encoding="utf-8")


def main():
    output = PROJECT / "plot"
    output.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 11,
        "axes.titlesize": 13,
        "axes.labelsize": 11,
        "legend.frameon": True,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
    })
    overview_figure(PROJECT, output)
    airfoil_figure(PROJECT, output)
    exact_summary = exact_load_figures(PROJECT, output)
    doe_figures(PROJECT, output, exact_summary)
    create_pptx(output)
    write_readme(output)
    print(f"Editable plot package written to {output}")


if __name__ == "__main__":
    main()

